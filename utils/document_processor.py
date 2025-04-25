import os
import logging
from typing import Optional

# Configure logging
logger = logging.getLogger("document-processor")

def extract_text_from_file(file_path: str) -> str:
    """
    Extract text from a document file.
    
    Supports:
    - PDF (.pdf)
    - Word (.docx)
    - Excel (.xlsx)
    - Text (.txt)
    
    Args:
        file_path: Path to document file
        
    Returns:
        Extracted text from document
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
        
    # Get file extension
    _, file_ext = os.path.splitext(file_path)
    file_ext = file_ext.lower()
    
    try:
        # Extract text based on file type
        if file_ext == '.pdf':
            return extract_text_from_pdf(file_path)
        elif file_ext == '.docx':
            return extract_text_from_docx(file_path)
        elif file_ext == '.xlsx':
            return extract_text_from_xlsx(file_path)
        elif file_ext == '.txt':
            return extract_text_from_txt(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_ext}")
    except Exception as e:
        logger.error(f"Error extracting text from {file_path}: {str(e)}")
        raise

def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF file"""
    try:
        from PyPDF2 import PdfReader
        
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    except ImportError:
        return f"[PDF Extraction Error: PyPDF2 not installed]"
    except Exception as e:
        return f"[PDF Extraction Error: {str(e)}]"

def extract_text_from_docx(file_path: str) -> str:
    """Extract text from DOCX file"""
    try:
        import docx
        
        doc = docx.Document(file_path)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        return text
    except ImportError:
        return f"[DOCX Extraction Error: python-docx not installed]"
    except Exception as e:
        return f"[DOCX Extraction Error: {str(e)}]"

def extract_text_from_xlsx(file_path: str) -> str:
    """Extract text from XLSX file"""
    try:
        import openpyxl
        
        workbook = openpyxl.load_workbook(file_path)
        text = ""
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                row_text = " ".join(str(cell.value) for cell in row if cell.value)
                if row_text:
                    text += row_text + "\n"
        return text
    except ImportError:
        return f"[XLSX Extraction Error: openpyxl not installed]"
    except Exception as e:
        return f"[XLSX Extraction Error: {str(e)}]"

def extract_text_from_txt(file_path: str) -> str:
    """Extract text from TXT file"""
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    except UnicodeDecodeError:
        # Try different encodings
        try:
            import chardet
            with open(file_path, 'rb') as f:
                raw_data = f.read()
            encoding = chardet.detect(raw_data)['encoding']
            with open(file_path, 'r', encoding=encoding) as f:
                return f.read()
        except ImportError:
            # If chardet is not available, try with common encodings
            for encoding in ['latin-1', 'cp1252', 'iso-8859-1']:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        return f.read()
                except UnicodeDecodeError:
                    continue
            return f"[TXT Extraction Error: Could not determine file encoding]"
    except Exception as e:
        return f"[TXT Extraction Error: {str(e)}]"